"""文本格式化和填充工具。"""

from pptx.enum.text import MSO_AUTO_SIZE


def copy_run_format(source_run, target_run) -> None:
    """复制 run 的格式属性。

    Args:
        source_run: 源 run 对象
        target_run: 目标 run 对象
    """
    try:
        if source_run.font.size:
            target_run.font.size = source_run.font.size
        if source_run.font.bold is not None:
            target_run.font.bold = source_run.font.bold
        if source_run.font.italic is not None:
            target_run.font.italic = source_run.font.italic
        if source_run.font.color.type is not None:
            target_run.font.color.rgb = source_run.font.color.rgb
        if source_run.font.name:
            target_run.font.name = source_run.font.name
    except Exception:
        pass


def copy_para_format(source_para, target_para) -> None:
    """复制段落的格式属性（缩进、对齐、行距等）。

    Args:
        source_para: 源段落对象
        target_para: 目标段落对象
    """
    try:
        target_para.alignment = source_para.alignment
        target_para.level = source_para.level
        if source_para.line_spacing:
            target_para.line_spacing = source_para.line_spacing
        if source_para.space_before:
            target_para.space_before = source_para.space_before
        if source_para.space_after:
            target_para.space_after = source_para.space_after
    except Exception:
        pass


def set_shape_text(shape, text) -> None:
    """填充文本时尽量保持原有格式。

    Args:
        shape: 形状对象
        text: 要填充的文本
    """
    if not shape.has_text_frame:
        return

    text = "" if text is None else str(text)
    lines = text.split("\n")
    tf = shape.text_frame

    # 先禁用自动调整大小，防止填充文本时形状被拉伸
    tf.auto_size = MSO_AUTO_SIZE.NONE

    if not tf.paragraphs:
        tf.add_paragraph()

    # 保存第一个段落和第一个 run 的格式作为模板
    template_para = tf.paragraphs[0] if tf.paragraphs else None
    template_run = (
        template_para.runs[0] if template_para and template_para.runs else None
    )

    for idx, line in enumerate(lines):
        if idx < len(tf.paragraphs):
            para = tf.paragraphs[idx]
        else:
            # 新建段落并复制格式
            para = tf.add_paragraph()
            if template_para:
                copy_para_format(template_para, para)

        if para.runs:
            para.runs[0].text = line
            for run in para.runs[1:]:
                run.text = ""
        else:
            # 没有 run 时创建一个并复制格式
            run = para.add_run()
            run.text = line
            if template_run:
                copy_run_format(template_run, run)

    # 清理多余段落
    for idx in range(len(lines), len(tf.paragraphs)):
        for run in tf.paragraphs[idx].runs:
            run.text = ""
